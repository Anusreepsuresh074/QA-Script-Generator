"""Build downloadable test execution reports (HTML, JSON, CSV, Excel)."""

from __future__ import annotations

import csv
import html
import io
import json
import math
from datetime import datetime, timezone
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def _escape(text: str) -> str:
    return html.escape(text or "")


def build_test_report_json(payload: dict[str, Any]) -> str:
    """Serialise a test-run payload to formatted JSON."""
    return json.dumps(payload, indent=2, ensure_ascii=False)


def build_test_report_csv(cases: Sequence[dict[str, Any]]) -> str:
    """Build a CSV of per-test results."""
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["outcome", "name", "classname", "duration_s", "message"])
    for case in cases:
        writer.writerow([
            case.get("outcome", ""),
            case.get("name", ""),
            case.get("classname", ""),
            case.get("duration", 0),
            (case.get("message") or "").replace("\n", " ").strip(),
        ])
    return buf.getvalue()


_OUTCOME_FILLS = {
    "passed": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
    "failed": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
    "error": PatternFill(start_color="FECACA", end_color="FECACA", fill_type="solid"),
    "skipped": PatternFill(start_color="F3E8FF", end_color="F3E8FF", fill_type="solid"),
}


def build_test_report_xlsx(payload: dict[str, Any]) -> bytes:
    """Build an Excel workbook with summary and per-test sheets."""
    ticket_id = str(payload.get("ticket_id", ""))
    total = int(payload.get("total", 0))
    passed = int(payload.get("passed", 0))
    failed = int(payload.get("failed", 0))
    errors = int(payload.get("errors", 0))
    skipped = int(payload.get("skipped", 0))
    failed_total = failed + errors
    pass_rate = _pct(passed, total)
    fail_rate = _pct(failed_total, total)

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"

    header_font = Font(bold=True, size=14)
    label_font = Font(bold=True)
    summary["A1"] = f"Test Report — {ticket_id}"
    summary["A1"].font = header_font
    summary["A2"] = f"Generated: {payload.get('generated_at', '')}"
    summary["A3"] = f"API base URL: {payload.get('base_url', '')}"
    summary["A4"] = f"Exit code: {payload.get('exit_code', 0)}"

    rows = [
        ("Verdict", "PASS" if failed_total == 0 and total > 0 else "FAIL" if total else "NO TESTS"),
        ("Total tests", total),
        ("Passed", passed),
        ("Failed", failed),
        ("Errors", errors),
        ("Skipped", skipped),
        ("Duration (s)", round(float(payload.get("duration", 0)), 2)),
        ("Pass rate (%)", pass_rate),
        ("Fail rate (%)", fail_rate),
    ]
    start = 6
    for idx, (label, value) in enumerate(rows, start=start):
        summary.cell(row=idx, column=1, value=label).font = label_font
        summary.cell(row=idx, column=2, value=value)

    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 28

    details = wb.create_sheet("Test Cases")
    headers = ["Outcome", "Test name", "Class", "Duration (s)", "Message"]
    for col, title in enumerate(headers, start=1):
        cell = details.cell(row=1, column=col, value=title)
        cell.font = label_font
        cell.fill = PatternFill(start_color="F4F7FB", end_color="F4F7FB", fill_type="solid")

    for row_idx, case in enumerate(payload.get("cases") or [], start=2):
        outcome = str(case.get("outcome", ""))
        values = [
            outcome,
            str(case.get("name", "")),
            str(case.get("classname", "")),
            float(case.get("duration", 0)),
            (case.get("message") or "").strip(),
        ]
        for col, value in enumerate(values, start=1):
            cell = details.cell(row=row_idx, column=col, value=value)
            if col == 1 and outcome in _OUTCOME_FILLS:
                cell.fill = _OUTCOME_FILLS[outcome]
            if col == 5:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    details.column_dimensions["A"].width = 12
    details.column_dimensions["B"].width = 42
    details.column_dimensions["C"].width = 18
    details.column_dimensions["D"].width = 14
    details.column_dimensions["E"].width = 60
    details.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pct(count: int, total: int) -> float:
    return round((count / total) * 100, 1) if total else 0.0


def _build_pie_chart_svg(passed: int, failed: int, errors: int, skipped: int) -> str:
    """Render an SVG pie chart for test outcomes (no external JS required)."""
    total = passed + failed + errors + skipped
    if total == 0:
        return (
            '<svg viewBox="0 0 200 200" class="pie-chart" aria-label="No test data">'
            '<circle cx="100" cy="100" r="80" fill="#e7ecf3"/>'
            '<text x="100" y="105" text-anchor="middle" fill="#5a6b82" font-size="14">No data</text>'
            "</svg>"
        )

    slices = [
        (passed, "#22c55e", "Passed"),
        (failed, "#ef4444", "Failed"),
        (errors, "#dc2626", "Errors"),
        (skipped, "#a855f7", "Skipped"),
    ]
    # Filter zero slices
    slices = [(v, c, l) for v, c, l in slices if v > 0]

    cx, cy, r = 100, 100, 80
    paths = []
    start_angle = -90  # start at top

    for value, color, _label in slices:
        sweep = (value / total) * 360
        if sweep >= 360:
            paths.append(f'<circle cx="{cx}" cy="{cy}" r="{r}" fill="{color}"/>')
            break
        end_angle = start_angle + sweep
        x1 = cx + r * math.cos(math.radians(start_angle))
        y1 = cy + r * math.sin(math.radians(start_angle))
        x2 = cx + r * math.cos(math.radians(end_angle))
        y2 = cy + r * math.sin(math.radians(end_angle))
        large_arc = 1 if sweep > 180 else 0
        paths.append(
            f'<path d="M{cx},{cy} L{x1:.2f},{y1:.2f} A{r},{r} 0 {large_arc},1 {x2:.2f},{y2:.2f} Z" fill="{color}"/>'
        )
        start_angle = end_angle

    return f'<svg viewBox="0 0 200 200" class="pie-chart" role="img" aria-label="Test results pie chart">{"".join(paths)}</svg>'


def _build_legend(passed: int, failed: int, errors: int, skipped: int, total: int) -> str:
    items = [
        ("Passed", passed, "#22c55e"),
        ("Failed", failed, "#ef4444"),
        ("Errors", errors, "#dc2626"),
        ("Skipped", skipped, "#a855f7"),
    ]
    rows = []
    for label, count, color in items:
        pct = _pct(count, total)
        rows.append(
            f"<li><span class='dot' style='background:{color}'></span>"
            f"<span class='legend-label'>{label}</span>"
            f"<span class='legend-count'>{count}</span>"
            f"<span class='legend-pct'>{pct}%</span></li>"
        )
    return f"<ul class='legend'>{''.join(rows)}</ul>"


def build_test_report_html(payload: dict[str, Any]) -> str:
    """Build a self-contained HTML test report with pie chart and outcome breakdown."""
    ticket_id = _escape(str(payload.get("ticket_id", "")))
    generated_at = _escape(str(payload.get("generated_at", "")))
    base_url = _escape(str(payload.get("base_url", "")))
    total = int(payload.get("total", 0))
    passed = int(payload.get("passed", 0))
    failed = int(payload.get("failed", 0))
    errors = int(payload.get("errors", 0))
    skipped = int(payload.get("skipped", 0))
    duration = float(payload.get("duration", 0))
    exit_code = int(payload.get("exit_code", 0))
    failed_total = failed + errors
    pass_rate = _pct(passed, total)
    fail_rate = _pct(failed_total, total)
    verdict = "PASS" if failed_total == 0 and total > 0 else "FAIL" if total else "NO TESTS"
    verdict_color = "#22c55e" if verdict == "PASS" else "#ef4444" if verdict == "FAIL" else "#f59e0b"
    pass_rate_color = "#22c55e" if pass_rate >= 80 else "#f59e0b" if pass_rate >= 50 else "#ef4444"

    pie_chart = _build_pie_chart_svg(passed, failed, errors, skipped)
    legend = _build_legend(passed, failed, errors, skipped, total)

    rows = []
    for case in payload.get("cases") or []:
        outcome = str(case.get("outcome", ""))
        color = {"passed": "#22c55e", "failed": "#ef4444", "error": "#dc2626", "skipped": "#a855f7"}.get(
            outcome, "#8b9cb3"
        )
        msg = case.get("message") or ""
        msg_html = f'<pre class="msg">{_escape(msg)}</pre>' if msg else ""
        rows.append(
            f"<tr>"
            f"<td><span class='badge' style='background:{color}'>{_escape(outcome)}</span></td>"
            f"<td><code>{_escape(str(case.get('name', '')))}</code>{msg_html}</td>"
            f"<td>{float(case.get('duration', 0)):.2f}s</td>"
            f"</tr>"
        )

    tbody = "\n".join(rows) or "<tr><td colspan='3'>No test cases recorded</td></tr>"
    raw_output = _escape(str(payload.get("output") or ""))

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8" />
  <title>Test Report — {ticket_id}</title>
  <style>
    body {{ font-family: Segoe UI, system-ui, sans-serif; margin: 2rem; color: #1a2332; max-width: 1100px; }}
    h1 {{ margin-bottom: 0.25rem; }}
    h2 {{ margin: 2rem 0 1rem; font-size: 1.1rem; }}
    .meta {{ color: #5a6b82; margin-bottom: 1.5rem; }}
    .summary {{ display: flex; gap: 1rem; flex-wrap: wrap; margin-bottom: 1.5rem; }}
    .card {{ border: 1px solid #d8e0ea; border-radius: 8px; padding: 1rem 1.25rem; min-width: 110px; }}
    .card span {{ color: #5a6b82; font-size: 0.85rem; }}
    .card strong {{ display: block; font-size: 1.5rem; margin-top: 0.2rem; }}
    .chart-section {{ display: flex; gap: 2rem; flex-wrap: wrap; align-items: center; margin: 1.5rem 0; padding: 1.5rem; background: #f8fafc; border-radius: 12px; border: 1px solid #e7ecf3; }}
    .pie-chart {{ width: 220px; height: 220px; flex-shrink: 0; }}
    .legend {{ list-style: none; padding: 0; margin: 0; min-width: 240px; }}
    .legend li {{ display: grid; grid-template-columns: 16px 1fr auto auto; gap: 0.6rem; align-items: center; padding: 0.45rem 0; border-bottom: 1px solid #e7ecf3; }}
    .legend li:last-child {{ border-bottom: none; }}
    .dot {{ width: 12px; height: 12px; border-radius: 50%; }}
    .legend-label {{ font-weight: 500; }}
    .legend-count {{ font-weight: 700; text-align: right; }}
    .legend-pct {{ color: #5a6b82; font-size: 0.9rem; min-width: 3rem; text-align: right; }}
    .breakdown {{ flex: 1; min-width: 220px; }}
    .breakdown p {{ margin: 0.4rem 0; line-height: 1.6; }}
    .pass-rate {{ font-size: 2rem; font-weight: 700; color: {pass_rate_color}; }}
    table {{ width: 100%; border-collapse: collapse; margin-top: 0.5rem; }}
    th, td {{ border-bottom: 1px solid #e7ecf3; padding: 0.65rem 0.5rem; text-align: left; vertical-align: top; }}
    th {{ background: #f4f7fb; }}
    .badge {{ color: #fff; padding: 0.15rem 0.5rem; border-radius: 999px; font-size: 0.75rem; text-transform: uppercase; }}
    code {{ font-family: ui-monospace, monospace; font-size: 0.9rem; }}
    pre.msg {{ margin: 0.5rem 0 0; padding: 0.5rem; background: #f8fafc; border-radius: 6px; white-space: pre-wrap; font-size: 0.8rem; }}
    details {{ margin-top: 2rem; }}
    @media print {{ body {{ margin: 1rem; }} .chart-section {{ break-inside: avoid; }} }}
  </style>
</head>
<body>
  <h1>Test Report — {ticket_id}</h1>
  <p class="meta">Generated {generated_at} · API base URL: {base_url} · Exit code {exit_code}</p>

  <div class="summary">
    <div class="card"><span>Verdict</span><strong style="color:{verdict_color}">{verdict}</strong></div>
    <div class="card"><span>Total tests</span><strong>{total}</strong></div>
    <div class="card"><span>Passed</span><strong style="color:#22c55e">{passed}</strong></div>
    <div class="card"><span>Failed</span><strong style="color:#ef4444">{failed}</strong></div>
    <div class="card"><span>Errors</span><strong style="color:#dc2626">{errors}</strong></div>
    <div class="card"><span>Skipped</span><strong style="color:#a855f7">{skipped}</strong></div>
    <div class="card"><span>Duration</span><strong>{duration:.2f}s</strong></div>
  </div>

  <h2>Results overview</h2>
  <div class="chart-section">
    {pie_chart}
    <div class="breakdown">
      <p class="pass-rate">{pass_rate}% pass rate</p>
      <p><strong>{passed}</strong> of <strong>{total}</strong> test cases passed.</p>
      <p><strong style="color:#ef4444">{failed_total}</strong> test case(s) failed or errored ({fail_rate}%).</p>
      {f'<p><strong style="color:#a855f7">{skipped}</strong> test case(s) skipped ({_pct(skipped, total)}%).</p>' if skipped else ''}
    </div>
    {legend}
  </div>

  <h2>Test case details</h2>
  <table>
    <thead><tr><th>Result</th><th>Test</th><th>Time</th></tr></thead>
    <tbody>{tbody}</tbody>
  </table>
  <details>
    <summary>Raw pytest output</summary>
    <pre>{raw_output}</pre>
  </details>
</body>
</html>"""


def report_filenames(ticket_id: str) -> dict[str, str]:
    """Standard report filenames for a ticket."""
    return {
        "html": f"{ticket_id}_test_report.html",
        "json": f"{ticket_id}_test_report.json",
        "csv": f"{ticket_id}_test_report.csv",
        "xlsx": f"{ticket_id}_test_report.xlsx",
    }


def is_report_filename(filename: str) -> bool:
    """Return True for generated test-report artifacts."""
    name = filename.lower()
    return name.endswith("_test_report.html") or name.endswith("_test_report.json") \
        or name.endswith("_test_report.csv") or name.endswith("_test_report.xlsx")


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")


def utc_stamp_for_path() -> str:
    """Filesystem-safe UTC timestamp for archived report folders."""
    return datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
